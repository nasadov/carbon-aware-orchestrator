"""
Build small repo-owned reference tables for water-related signals.

The generated CSVs are derived from official upstream datasets:
- AWARE2.0 country non-agricultural characterization factors
- WRI purchased-electricity embedded water factors

This helper is intentionally offline at runtime; the scheduler only reads the
generated CSVs.
"""
from __future__ import annotations

import csv
from pathlib import Path
from tempfile import TemporaryDirectory
from typing import Dict, Iterable

import openpyxl
import requests


THIS_DIR = Path(__file__).resolve().parent

AWARE_URL = "https://zenodo.org/records/16332127/files/AWARE20_Countries_and_Regions.xlsx?download=1"
WRI_URL = "https://files.wri.org/d8/s3fs-public/guidance-calculating-water-use-embedded-purchased-electricity-appendices-1-to-7.xlsx"
GALLON_TO_LITER = 3.785411784


def _download(url: str, dest: Path) -> None:
    response = requests.get(url, timeout=120)
    response.raise_for_status()
    dest.write_bytes(response.content)


def _safe_float(value) -> str:
    if value in (None, "", "NotDefined"):
        return ""
    return str(float(value))


def _build_aware_country_map(workbook_path: Path) -> Dict[str, str]:
    wb = openpyxl.load_workbook(workbook_path, read_only=True, data_only=True)
    ws = wb["CFs_nonagri"]
    mapping: Dict[str, str] = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        country_name = row[4]
        country_code = row[6]
        if country_name and country_code:
            mapping[str(country_name)] = str(country_code).upper()
    return mapping


def _write_aware_country_factors(workbook_path: Path, output_path: Path) -> None:
    wb = openpyxl.load_workbook(workbook_path, read_only=True, data_only=True)
    ws = wb["CFs_nonagri"]
    month_labels = ["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"]

    with output_path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(
            [
                "country_name",
                "country_code",
                "annual_cf",
                *[f"{month}_cf" for month in month_labels],
                "source_dataset",
                "source_url",
            ]
        )
        for row in ws.iter_rows(min_row=2, values_only=True):
            country_name = row[4]
            country_code = row[6]
            if not country_name or not country_code:
                continue
            writer.writerow(
                [
                    country_name,
                    str(country_code).upper(),
                    _safe_float(row[8]),
                    *[_safe_float(row[idx]) for idx in range(9, 21)],
                    "AWARE2.0 CFs_nonagri",
                    AWARE_URL,
                ]
            )


def _iter_wri_country_rows(workbook_path: Path) -> Iterable[tuple]:
    wb = openpyxl.load_workbook(workbook_path, read_only=True, data_only=True)
    ws = wb["Appendix 1 & 2"]
    for row in ws.iter_rows(min_row=3, values_only=True):
        country_name = row[4]
        withdrawal = row[5]
        consumption = row[6]
        if country_name and withdrawal is not None and consumption is not None:
            yield country_name, withdrawal, consumption


def _write_wri_country_factors(workbook_path: Path, aware_country_codes: Dict[str, str], output_path: Path) -> None:
    with output_path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(
            [
                "country_name",
                "country_code",
                "withdrawal_gal_per_kwh",
                "consumption_gal_per_kwh",
                "withdrawal_l_per_kwh",
                "consumption_l_per_kwh",
                "source_dataset",
                "source_url",
            ]
        )
        for country_name, withdrawal, consumption in _iter_wri_country_rows(workbook_path):
            country_code = aware_country_codes.get(str(country_name))
            if not country_code:
                continue
            writer.writerow(
                [
                    country_name,
                    country_code,
                    float(withdrawal),
                    float(consumption),
                    float(withdrawal) * GALLON_TO_LITER,
                    float(consumption) * GALLON_TO_LITER,
                    "WRI purchased-electricity appendix country factors",
                    WRI_URL,
                ]
            )


def main() -> None:
    with TemporaryDirectory() as tmpdir:
        tmpdir_path = Path(tmpdir)
        aware_xlsx = tmpdir_path / "aware20.xlsx"
        wri_xlsx = tmpdir_path / "wri_embedded_electricity_water.xlsx"

        _download(AWARE_URL, aware_xlsx)
        _download(WRI_URL, wri_xlsx)

        aware_country_codes = _build_aware_country_map(aware_xlsx)
        _write_aware_country_factors(aware_xlsx, THIS_DIR / "aware20_country_nonagri_factors.csv")
        _write_wri_country_factors(wri_xlsx, aware_country_codes, THIS_DIR / "wri_country_grid_water_factors.csv")


if __name__ == "__main__":
    main()
